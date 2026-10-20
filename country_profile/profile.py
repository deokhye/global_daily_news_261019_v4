"""최신 프로필 생성: output/country_profile.json, output/country_profile.xlsx.

항목 순서(대시보드 '국가 기본 프로필' 카드 기준):
  수도 → 총인구 → 연간 GDP 규모 → 인플레이션 → 실업률 → 최저임금 → 한국타이어 주요 거점
표기 형식(별도 양식이 없을 때의 기본값):
  총인구 51.7백만 명 / GDP 1,872.3십억 USD / 인플레이션·실업률 2.9% (2026.09)
"""
from __future__ import annotations

import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import storage
from .config import KST, Country, Paths

REVIEW_KEYWORDS = ("검토 필요", "미갱신", "대체 출처", "경고")
LABELS = {"population": "총인구", "gdp": "연간 GDP 규모", "inflation": "인플레이션",
          "unemployment": "실업률", "min_wage": "최저임금"}


def fmt_period(p: str) -> str:
    p = p or ""
    if len(p) == 7 and p[4] == "-" and p[5:].isdigit():
        return f"{p[:4]}.{p[5:]}"
    return p.replace("-", " ")


def fmt_amount(v: float) -> str:
    s = f"{v:,.2f}"
    return s[:-3] if s.endswith(".00") else s


def _latest_annual(df, iso3: str, indicator: str):
    sub = df[(df.iso3 == iso3) & (df.indicator == indicator)] if not df.empty else df
    if sub.empty:
        return None
    return sub.assign(_y=sub.base_year.astype(int)).sort_values("_y").iloc[-1]


def _latest_monthly(df, iso3: str, indicator: str):
    sub = df[(df.iso3 == iso3) & (df.indicator == indicator)] if not df.empty else df
    if sub.empty:
        return None
    return sub.sort_values("snapshot_date").iloc[-1]


def _display(indicator: str, row) -> str:
    if row is None:
        return "-"
    v = storage.to_float(row["value"])
    if indicator == "min_wage":
        if str(row["status"]).startswith("법정 최저임금 없음"):
            note = row.get("note", "")
            return f"법정 최저임금 없음 ({note.split('—')[-1].strip()})" if note else "법정 최저임금 없음"
        if v is None:
            return "-"
        parts = [f"{row['currency']} {fmt_amount(v)} / {row['unit'] or '단위 미상'}".strip()]
        extra = [x for x in (row.get("region", ""), f"{row['effective_date']} 시행" if row["effective_date"] else "") if x]
        return parts[0] + (f" ({', '.join(extra)})" if extra else "")
    if v is None:
        return "-"
    if indicator == "population":
        return f"{v:,.1f}백만 명"
    if indicator == "gdp":
        return f"{v:,.1f}십억 USD"
    return f"{v:.1f}% ({fmt_period(row['ref_period'])})" if row.get("ref_period") else f"{v:.1f}%"


def build_profile(countries: list[Country], paths: Paths) -> dict:
    annual = storage.load(paths.annual_history, storage.ANNUAL_COLUMNS)
    monthly = storage.load(paths.monthly_history, storage.MONTHLY_COLUMNS)
    out, review = [], []

    for c in countries:
        rows = {ind: _latest_annual(annual, c.iso3, ind) for ind in ("population", "gdp", "min_wage")}
        rows.update({ind: _latest_monthly(monthly, c.iso3, ind) for ind in ("inflation", "unemployment")})
        item = {
            "region": c.region, "country_ko": c.name_ko, "country_en": c.name_en, "iso3": c.iso3,
            "iso2": c.iso2,
            "capital": c.capital,
            "population": _display("population", rows["population"]),
            "gdp": _display("gdp", rows["gdp"]),
            "inflation": _display("inflation", rows["inflation"]),
            "unemployment": _display("unemployment", rows["unemployment"]),
            "min_wage": _display("min_wage", rows["min_wage"]),
            "sites": c.sites,
            "details": {},
        }
        for ind, row in rows.items():
            if row is None:
                item["details"][ind] = {"status": "데이터 없음"}
                review.append({"iso3": c.iso3, "국가": c.name_ko, "항목": LABELS[ind], "상태": "데이터 없음",
                               "값": "-", "출처": "", "출처 URL": ""})
                continue
            d = {k: row[k] for k in row.index if k in storage.ANNUAL_COLUMNS + storage.MONTHLY_COLUMNS}
            item["details"][ind] = d
            status = str(row["status"])
            if any(k in status for k in REVIEW_KEYWORDS):
                review.append({"iso3": c.iso3, "국가": c.name_ko, "항목": LABELS[ind], "상태": status,
                               "값": item[ind], "출처": row.get("source", "Trading Economics"),
                               "출처 URL": row.get("source_url", "")})
        out.append(item)

    profile = {"generated_at": datetime.now(KST).isoformat(timespec="seconds"), "countries": out, "review": review}
    paths.output.mkdir(parents=True, exist_ok=True)
    with open(paths.output / "country_profile.json", "w", encoding="utf-8") as f:
        json.dump(profile, f, ensure_ascii=False, indent=2)
    _write_xlsx(profile, paths.output / "country_profile.xlsx")
    return profile


def _write_xlsx(profile: dict, path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "국가 기본 프로필"
    head = ["권역", "국가", "영문명", "ISO3", "수도", "총인구", "연간 GDP 규모", "인플레이션 (CPI YoY)",
            "실업률", "최저임금", "한국타이어 주요 거점"]
    ws.append(head)
    for c in profile["countries"]:
        ws.append([c["region"], c["country_ko"], c["country_en"], c["iso3"], c["capital"], c["population"],
                   c["gdp"], c["inflation"], c["unemployment"], c["min_wage"], "\n".join(c["sites"])])
    widths = [8, 12, 22, 7, 18, 16, 18, 20, 16, 44, 60]
    _style(ws, widths)

    rv = wb.create_sheet("검토 필요")
    rhead = ["iso3", "국가", "항목", "상태", "값", "출처", "출처 URL"]
    rv.append(rhead)
    for r in profile["review"]:
        rv.append([r[h] for h in rhead])
    _style(rv, [7, 12, 14, 60, 36, 28, 50])

    meta = wb.create_sheet("정보")
    meta.append(["생성 시각(KST)", profile["generated_at"]])
    meta.append(["상태 표기", "정상 / 갱신 없음 / 대체 출처(World Bank) / 검토 필요 / 미갱신 / 법정 최저임금 없음 / 경고: 급변"])
    wb.save(path)


def _style(ws, widths):
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="F27A18")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2" if ws.max_column > 7 else "A2"
